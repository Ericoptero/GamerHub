"""Importador genérico de planilha (XLSX ou CSV) em duas etapas: prévia com mapeamento sugerido e commit."""

import csv
import io
import secrets
from pathlib import Path

from openpyxl import load_workbook
from pydantic import BaseModel
from rapidfuzz import fuzz
from sqlalchemy.orm import Session

from app.config import get_settings
from app.domain import fold
from app.services.catalog import split_stores
from app.services.importer import IMPORT_MODES, GameRow, ImportReport, upsert_rows
from app.services.normalize import (
    clean_text,
    parse_date,
    parse_int,
    parse_number,
    parse_priority,
    parse_status,
)

# campo -> (rótulo, sinônimos de cabeçalho)
FIELDS: dict[str, tuple[str, list[str]]] = {
    "name": ("Nome", ["nome", "jogo", "titulo", "game", "name", "nome do jogo"]),
    "platform": ("Plataforma", ["plataforma", "console", "sistema", "platform"]),
    "genre": ("Gênero", ["genero", "categoria", "genre", "tipo"]),
    "status": ("Lista / status", ["lista", "status", "situacao", "estado"]),
    "priority": ("Prioridade", ["prioridade", "priority", "prio"]),
    "hltb_hours": ("HLTB (horas)", ["hltb", "hltb (h)", "how long to beat", "horas hltb", "tempo hltb", "duracao"]),
    "stores": ("Lojas", ["lojas", "loja", "origem", "store", "stores", "onde tenho"]),
    "date": ("Data", ["data", "finalizado em", "zerado em", "data de finalizacao", "concluido em", "lancamento"]),
    "played_hours": ("Horas jogadas", ["horas jogadas", "tempo jogado", "tempo", "horas", "played", "meu tempo"]),
    "user_score": ("Nota", ["nota", "minha nota", "score", "avaliacao", "sua nota"]),
    "review": ("Review", ["review", "comentario", "comentarios", "opiniao", "analise", "o que achei"]),
    "story": ("História", ["historia", "resumo", "enredo", "sinopse", "resumo da historia"]),
    "times_completed": ("Vezes zerado", ["zerado (vezes)", "vezes zerado", "zerado", "vezes", "zeradas"]),
    "metacritic": ("Metacritic", ["metacritic", "mc", "metascore", "nota metacritic"]),
}


class SheetInfo(BaseModel):
    name: str
    headers: list[str]
    sample: list[list[str]]
    row_count: int
    mapping: dict[str, int]
    status_hint: str | None


class SheetPreview(BaseModel):
    token: str
    filename: str
    fields: dict[str, str]
    sheets: list[SheetInfo]


class SheetChoice(BaseModel):
    name: str
    include: bool = True
    mapping: dict[str, int]
    status_override: str | None = None


class SheetCommit(BaseModel):
    token: str
    sheets: list[SheetChoice]
    mode: str = "novos"
    dry_run: bool = False


def _cell_str(v: object) -> str:
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%d/%m/%Y")  # type: ignore[union-attr]
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _read_csv(data: bytes) -> list[list[object]]:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Não consegui ler o CSV (encoding desconhecido)")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if text.count(";") > text.count(",") else ","
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]


def read_workbook(path: Path) -> dict[str, list[list[object]]]:
    """{nome da aba: linhas}. CSV vira uma aba só com o nome do arquivo."""
    if path.suffix.lower() == ".csv":
        return {path.stem.split("_", 1)[-1]: _read_csv(path.read_bytes())}
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001  (openpyxl levanta vários tipos)
        raise ValueError(f"Não consegui abrir a planilha: {e}") from None
    sheets: dict[str, list[list[object]]] = {}
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        sheets[ws.title] = rows
    wb.close()
    return sheets


def _split_header(rows: list[list[object]]) -> tuple[list[str], list[list[object]]]:
    for i, row in enumerate(rows):
        if any(c not in (None, "") for c in row):
            headers = [_cell_str(c).strip() for c in row]
            data = [r for r in rows[i + 1 :] if any(c not in (None, "") for c in r)]
            return headers, data
    return [], []


def suggest_mapping(headers: list[str]) -> dict[str, int]:
    """Casa cabeçalhos com campos: igualdade com sinônimo primeiro, depois similaridade."""
    folded = [fold(h) for h in headers]
    scores: list[tuple[float, str, int]] = []
    for field, (_label, synonyms) in FIELDS.items():
        for idx, h in enumerate(folded):
            if not h:
                continue
            if h in synonyms:
                best = 100.0 + (len(h) / 100)
            else:
                best = max(fuzz.ratio(h, s) for s in synonyms)
                if best < 85:
                    continue
            scores.append((best, field, idx))
    mapping: dict[str, int] = {}
    used: set[int] = set()
    for _score, field, idx in sorted(scores, reverse=True):
        if field not in mapping and idx not in used:
            mapping[field] = idx
            used.add(idx)
    return mapping


def preview(filename: str, data: bytes) -> SheetPreview:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".xlsm", ".csv"}:
        raise ValueError("Envie um arquivo .xlsx ou .csv")
    token = secrets.token_hex(8)
    imports = get_settings().imports_dir
    imports.mkdir(parents=True, exist_ok=True)
    safe_stem = "".join(c for c in Path(filename).stem if c.isalnum() or c in " -_")[:60] or "planilha"
    path = imports / f"{token}_{safe_stem}{suffix}"
    path.write_bytes(data)

    sheets: list[SheetInfo] = []
    for name, rows in read_workbook(path).items():
        headers, body = _split_header(rows)
        if not headers:
            continue
        sheets.append(
            SheetInfo(
                name=name,
                headers=headers,
                sample=[[_cell_str(c) for c in r] for r in body[:10]],
                row_count=len(body),
                mapping=suggest_mapping(headers),
                status_hint=parse_status(name),
            )
        )
    if not sheets:
        raise ValueError("A planilha está vazia")
    return SheetPreview(
        token=token, filename=filename, fields={k: v[0] for k, v in FIELDS.items()}, sheets=sheets
    )


def _find_upload(token: str) -> Path:
    if not token.isalnum():
        raise ValueError("token inválido")
    matches = list(get_settings().imports_dir.glob(f"{token}_*"))
    if not matches:
        raise FileNotFoundError("Arquivo da prévia não encontrado; envie de novo")
    return matches[0]


def row_from_cells(cells: list[object], mapping: dict[str, int], status: str | None, line: str) -> GameRow:
    def get(field: str) -> object:
        idx = mapping.get(field)
        if idx is None or idx >= len(cells):
            return None
        v = cells[idx]
        return None if v == "" else v

    d_value, d_label = parse_date(get("date"))
    played = parse_number(get("played_hours"))
    return GameRow(
        name=(clean_text(get("name")) or ""),
        platform=clean_text(get("platform")),
        genre=clean_text(get("genre")),
        status=parse_status(get("status")) or status,
        priority=parse_priority(get("priority")),
        hltb_hours=parse_number(get("hltb_hours")),
        stores=split_stores(clean_text(get("stores"))),
        date_value=d_value,
        date_label=d_label,
        played_hours=played if played else None,
        user_score=parse_int(get("user_score")),
        review=clean_text(get("review")),
        story=clean_text(get("story")),
        times_completed=parse_int(get("times_completed")),
        metacritic=parse_int(get("metacritic")),
        line=line,
    )


def commit(db: Session, data: SheetCommit) -> ImportReport:
    if data.mode not in IMPORT_MODES:
        raise ValueError(f"modo deve ser um de {IMPORT_MODES}")
    path = _find_upload(data.token)
    book = read_workbook(path)
    rows: list[GameRow] = []
    for choice in data.sheets:
        if not choice.include or choice.name not in book:
            continue
        if "name" not in choice.mapping:
            raise ValueError(f"Aba {choice.name!r}: mapeie a coluna do nome do jogo")
        status = choice.status_override or parse_status(choice.name)
        _headers, body = _split_header(book[choice.name])
        for i, cells in enumerate(body):
            rows.append(row_from_cells(cells, choice.mapping, status, f"{choice.name} linha {i + 2}"))
    report = upsert_rows(db, rows, data.mode, dry_run=data.dry_run)
    if not data.dry_run:
        path.unlink(missing_ok=True)
    return report
